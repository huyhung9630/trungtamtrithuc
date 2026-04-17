"""Document parser — 3-tier: Docling → Claude Vision → pdfplumber.

Flow:
  PDF → Docling → Markdown (always return, pipeline handles tables)
    - Docling quality bad → Vision per-page (context mode) → fallback pdfplumber
  DOCX → Docling → Markdown, fallback python-docx
  XLSX → openpyxl (structured text)
  TXT/MD → read directly
"""
from __future__ import annotations

import base64
import hashlib
import logging
import mimetypes
import re
from datetime import datetime, timezone
from pathlib import Path
from typing import Union

logger = logging.getLogger(__name__)

_docling_converter = None

MIN_TEXT_CHARS = 50

VISION_TO_MARKDOWN_PROMPT = (
    "Hãy chuyển nội dung trang tài liệu này thành Markdown.\n"
    "Yêu cầu:\n"
    "- Nếu có bảng: chuyển thành Markdown table (| cột 1 | cột 2 |)\n"
    "- Nếu có tiêu đề: dùng ## heading\n"
    "- Nếu có danh sách: dùng - bullet\n"
    "- QUAN TRỌNG: Giữ nguyên chính xác tiếng Việt có dấu.\n"
    "  Ví dụ: 'Chống cháy' KHÔNG PHẢI 'Chông cháy', 'Cửa' KHÔNG PHẢI 'Của'\n"
    "- Chỉ trả về Markdown, không giải thích thêm."
)

VISION_TO_CONTEXT_PROMPT = (
    "Đọc trang tài liệu này và viết lại NỘI DUNG thành plain text có cấu trúc.\n"
    "Yêu cầu:\n"
    "- Dòng đầu: tiêu đề/mô tả ngắn về trang này.\n"
    "- Nếu có bảng: mỗi dòng dữ liệu viết thành 1 dòng text riêng.\n"
    "  VD: 'Video 1: Giới thiệu BKVN. Tình trạng: Lên kịch bản.'\n"
    "- Nếu ô có màu mang ý nghĩa (xanh=hoàn thành, cam=đang làm, đỏ=quan trọng):\n"
    "  → Ghi gắn với dữ liệu trên cùng dòng.\n"
    "  VD: 'Tất niên, du xuân (T1-T2): đã hoàn thành.'\n"
    "- Dữ liệu lặp: ghi 'Tháng 3, 4: lặp lại nội dung Tháng 2.'\n"
    "- Bỏ qua cột/ô trống.\n"
    "- Giữ chính xác tên riêng, số liệu, tiếng Việt có dấu.\n"
    "  VD: 'Chống cháy' KHÔNG PHẢI 'Chông cháy'\n"
    "- KHÔNG dùng Markdown (**, ##, `, |, ---).\n"
    "- KHÔNG mô tả format, font, kiểu bảng.\n"
    "- Chỉ trả về plain text."
)

# Common OCR/Vision typos
_TYPO_FIXES: list[tuple[str, str]] = [
    ("Chông Cháy", "Chống Cháy"), ("Chông cháy", "Chống cháy"),
    ("chông cháy", "chống cháy"), ("Chông Chảy", "Chống Cháy"),
    ("Chông chảy", "Chống cháy"), ("chông chảy", "chống cháy"),
    ("Của chống", "Cửa chống"), ("Của Chống", "Cửa Chống"),
    ("của chống", "cửa chống"), ("Của chông", "Cửa chống"),
    ("Của Chông", "Cửa Chống"), ("của chông", "cửa chống"),
    ("Dã quay", "Đã quay"), ("Dã dạng", "Đã đăng"),
    ("Dạng edit", "Đang edit"), ("Tính trạng", "Tình trạng"),
    ("KỀ HOẠCH", "KẾ HOẠCH"), ("Kề hoạch", "Kế hoạch"),
    ("kề hoạch", "kế hoạch"), ("ban nhạt định", "bạn nhất định"),
    ("Ban nhạt định", "Bạn nhất định"), ("phải ban nhạt", "phải bạn nhất"),
]


def _sha256(path: Path) -> str:
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(65536), b""):
            h.update(chunk)
    return h.hexdigest()[:16]


def _fix_common_typos(text: str) -> str:
    for wrong, correct in _TYPO_FIXES:
        text = text.replace(wrong, correct)
    return text


def _clean_markdown(md: str) -> str:
    md = re.sub(r"<!--.*?-->", "", md, flags=re.DOTALL)
    md = re.sub(r"-{3,}\s*page\s*\d*\s*-{3,}", "", md, flags=re.IGNORECASE)
    md = re.sub(r"\n{3,}", "\n\n", md)
    lines = [line.rstrip() for line in md.split("\n")]
    return "\n".join(lines).strip()


def _fix_joined_words(text: str) -> str:
    text = re.sub(
        r"([a-zàáảãạăắằẳẵặâấầẩẫậèéẻẽẹêếềểễệìíỉĩịòóỏõọôốồổỗộơớờởỡợùúủũụưứừửữựỳýỷỹỵđ])"
        r"([A-ZÀÁẢÃẠĂẮẰẲẴẶÂẤẦẨẪẬÈÉẺẼẸÊẾỀỂỄỆÌÍỈĨỊÒÓỎÕỌÔỐỒỔỖỘƠỚỜỞỠỢÙÚỦŨỤƯỨỪỬỮỰỲÝỶỸỴĐ])",
        r"\1 \2", text,
    )
    return text


def _count_meaningful_chars(text: str) -> int:
    return len(re.sub(r"[|#\-=_*`~>\s]", "", text))


def _check_docling_quality(md: str) -> tuple[bool, str]:
    if not md or not md.strip():
        return False, "empty output"
    lines = md.strip().split("\n")
    non_empty = [l for l in lines if l.strip()]
    if len(non_empty) <= 2 and md.count("|") > 20:
        return False, f"broken table: {md.count('|')} pipes on {len(non_empty)} lines"
    cells = md[:200].split("|")
    if len(cells) >= 3:
        first_cell = cells[1].strip() if len(cells) > 1 else ""
        if first_cell and len(first_cell) > 5 and md.count(first_cell) >= 5:
            return False, f"repeated header x{md.count(first_cell)}"
    if len(md) > 500 and len(non_empty) < 3:
        return False, f"no line breaks: {len(md)} chars in {len(non_empty)} lines"
    if len(md) > 500:
        paragraphs = [p.strip() for p in re.split(r"\n{2,}", md) if len(p.strip()) > 100]
        if len(paragraphs) >= 4:
            unique = set(paragraphs)
            dup_ratio = 1 - (len(unique) / len(paragraphs))
            if dup_ratio > 0.4:
                return False, f"duplicate blocks: {dup_ratio:.0%}"
    return True, "ok"


# --- Tier 1: Docling ---

def _get_docling_converter():
    global _docling_converter
    if _docling_converter is not None:
        return _docling_converter
    try:
        from docling.document_converter import DocumentConverter
        _docling_converter = DocumentConverter()
        logger.info("Docling converter loaded")
    except Exception as exc:
        logger.warning("Docling not available: %s", exc)
    return _docling_converter


def _docling_parse(path: Path) -> str:
    converter = _get_docling_converter()
    if converter is None:
        return ""
    try:
        result = converter.convert(source=str(path))
        md = result.document.export_to_markdown()
        md = _clean_markdown(md)
        md = _fix_joined_words(md)
        is_good, reason = _check_docling_quality(md)
        if not is_good:
            logger.warning("Docling REJECTED for %s: %s", path.name, reason)
            return ""
        logger.info("Docling OK for %s: %d chars", path.name, len(md))
        return md
    except Exception:
        logger.exception("Docling failed for %s", path.name)
        return ""


# --- Tier 2: Claude Vision ---

def _get_anthropic_client():
    from app.config import ANTHROPIC_API_KEY
    if not ANTHROPIC_API_KEY:
        return None
    try:
        import anthropic
        return anthropic.Anthropic(api_key=ANTHROPIC_API_KEY)
    except Exception:
        return None


def _pdf_page_to_base64(path: Path, page_num: int) -> str | None:
    try:
        from pdf2image import convert_from_path
        images = convert_from_path(str(path), dpi=200, first_page=page_num, last_page=page_num)
        if not images:
            return None
        import io
        buf = io.BytesIO()
        images[0].save(buf, format="PNG")
        b64 = base64.b64encode(buf.getvalue()).decode("utf-8")
        del images
        return b64
    except Exception:
        logger.exception("Failed to convert page %d to image", page_num)
        return None


def _vision_call(client, image_b64: str, model: str, prompt: str) -> str:
    try:
        response = client.messages.create(
            model=model, max_tokens=4096, temperature=0.1,
            messages=[{"role": "user", "content": [
                {"type": "image", "source": {"type": "base64", "media_type": "image/png", "data": image_b64}},
                {"type": "text", "text": prompt},
            ]}],
        )
        result = response.content[0].text.strip()
        return _fix_common_typos(result)
    except Exception:
        logger.exception("Claude Vision call failed")
        return ""


def _vision_parse_pdf(path: Path, to_context: bool = False) -> list[dict]:
    from app.config import CLAUDE_HAIKU_MODEL
    client = _get_anthropic_client()
    if client is None:
        return []
    num_pages = _count_pdf_pages(path)
    prompt = VISION_TO_CONTEXT_PROMPT if to_context else VISION_TO_MARKDOWN_PROMPT
    logger.info("Vision parsing %s: %d pages, context=%s", path.name, num_pages, to_context)
    pages: list[dict] = []
    for page_num in range(1, num_pages + 1):
        image_b64 = _pdf_page_to_base64(path, page_num)
        if not image_b64:
            pages.append({"page": page_num, "text": ""})
            continue
        text = _vision_call(client, image_b64, CLAUDE_HAIKU_MODEL, prompt)
        text = _clean_markdown(text)
        logger.info("Vision page %d: %d chars", page_num, len(text))
        pages.append({"page": page_num, "text": text})
    return pages


# --- Tier 3: pdfplumber ---

def _pdfplumber_parse(path: Path) -> list[dict]:
    try:
        import pdfplumber
    except ImportError:
        return []
    pages: list[dict] = []
    try:
        with pdfplumber.open(path) as pdf:
            for i, page in enumerate(pdf.pages, start=1):
                try:
                    text = (page.extract_text() or "").strip()
                except Exception:
                    text = ""
                pages.append({"page": i, "text": text})
    except Exception:
        logger.exception("pdfplumber cannot open %s", path.name)
    return pages


def _count_pdf_pages(path: Path) -> int:
    try:
        import pdfplumber
        with pdfplumber.open(path) as pdf:
            return len(pdf.pages)
    except Exception:
        return 1


# --- Main PDF parser ---

def parse_pdf(path: Path) -> list[dict]:
    # Tier 1: Docling — always return if quality OK
    markdown = _docling_parse(path)
    if markdown and _count_meaningful_chars(markdown) >= MIN_TEXT_CHARS:
        return [{"page": 1, "text": markdown}]

    # Tier 2: Vision (Docling failed → context mode)
    logger.info("PDF %s: Docling failed, trying Vision", path.name)
    vision_pages = _vision_parse_pdf(path, to_context=True)
    non_empty = [p for p in vision_pages if p["text"].strip()]
    if non_empty:
        return vision_pages

    # Tier 3: pdfplumber
    logger.warning("PDF %s: Vision failed, using pdfplumber", path.name)
    pages = _pdfplumber_parse(path)
    return pages or [{"page": 1, "text": ""}]


# --- XLSX (kept from teammate's code) ---

def parse_xlsx(path: Path) -> str:
    from openpyxl import load_workbook
    wb = load_workbook(str(path), data_only=True)
    parts: list[str] = []
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            continue
        header_idx = 0
        headers: list[str] = []
        for i, row in enumerate(rows):
            non_empty = [c for c in row if c is not None and str(c).strip()]
            if len(non_empty) >= 2:
                headers = [str(c).strip() if c else f"Col{j+1}" for j, c in enumerate(row)]
                header_idx = i
                break
        if len(wb.sheetnames) > 1:
            parts.append(f"# Sheet: {sheet_name}")
        if not headers:
            for row in rows:
                line = " | ".join(str(c).strip() for c in row if c is not None and str(c).strip())
                if line:
                    parts.append(line)
            continue
        for row in rows[header_idx + 1:]:
            cells = [str(c).strip() if c is not None else "" for c in row]
            if not any(cells):
                continue
            row_parts: list[str] = []
            for header, cell in zip(headers, cells):
                if cell:
                    row_parts.append(f"{header}: {cell}")
            if row_parts:
                parts.append("\n".join(row_parts))
    return "\n\n".join(parts)


# --- DOCX, TXT, MD ---

def parse_docx(path: Path) -> str:
    converter = _get_docling_converter()
    if converter is not None:
        try:
            result = converter.convert(source=str(path))
            md = result.document.export_to_markdown()
            md = _clean_markdown(md)
            md = _fix_joined_words(md)
            if _count_meaningful_chars(md) >= MIN_TEXT_CHARS:
                return md
        except Exception:
            logger.exception("Docling DOCX failed for %s", path.name)
    try:
        from docx import Document
        doc = Document(str(path))
        parts = []
        for para in doc.paragraphs:
            text = para.text.strip()
            if not text:
                continue
            style = para.style.name if para.style else ""
            if style.startswith("Heading"):
                level = "".join(filter(str.isdigit, style)) or "1"
                parts.append("#" * int(level) + " " + text)
            else:
                parts.append(text)
        return "\n\n".join(parts)
    except Exception:
        logger.exception("python-docx failed for %s", path.name)
        return ""


def parse_txt(path: Path) -> str:
    for encoding in ("utf-8", "utf-8-sig", "cp1252", "latin-1"):
        try:
            return path.read_text(encoding=encoding)
        except (UnicodeDecodeError, ValueError):
            continue
    return path.read_text(encoding="utf-8", errors="replace")


def parse_md(path: Path) -> str:
    return parse_txt(path)


# --- Main entry ---

def parse(path: Union[str, Path]) -> dict:
    path = Path(path)
    if not path.exists():
        logger.error("File not found: %s", path)
        return {"doc_id": "0" * 16, "source": path.name, "content": "",
                "mime": "application/octet-stream", "uploaded_at": datetime.now(timezone.utc).isoformat()}
    if path.stat().st_size == 0:
        logger.warning("Empty file: %s", path.name)
        return {"doc_id": _sha256(path), "source": path.name, "content": "",
                "mime": "application/octet-stream", "uploaded_at": datetime.now(timezone.utc).isoformat()}

    suffix = path.suffix.lower()
    doc_id = _sha256(path)
    uploaded_at = datetime.now(timezone.utc).isoformat()

    if suffix == ".pdf":
        content = parse_pdf(path)
        mime = "application/pdf"
    elif suffix in (".docx", ".doc"):
        content = parse_docx(path)
        mime = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    elif suffix == ".xlsx":
        content = parse_xlsx(path)
        mime = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    elif suffix == ".md":
        content = parse_md(path)
        mime = "text/markdown"
    else:
        content = parse_txt(path)
        mime = mimetypes.guess_type(path.name)[0] or "text/plain"

    return {"doc_id": doc_id, "source": path.name, "content": content,
            "mime": mime, "uploaded_at": uploaded_at}
